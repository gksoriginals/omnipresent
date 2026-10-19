"""Synthetic, clearly labeled fixtures for local extraction checks."""
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
from pptx import Presentation
from openpyxl import Workbook
out = Path('output/probe/material-fixtures')
out.mkdir(parents=True, exist_ok=True)
text = 'Synthetic solar material. Panels generate electricity. Batteries store energy.'
with ZipFile(out / 'notes.docx', 'w') as z:
    z.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    z.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    z.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>'+text+'</w:t></w:r></w:p></w:body></w:document>')
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = 'Synthetic solar material'
slide.placeholders[1].text = text
presentation.save(out / 'notes.pptx')
wb = Workbook(); ws = wb.active
ws.append(['Period', 'Energy']); ws.append(['Morning', 10]); ws.append(['Noon', 20]); wb.save(out / 'data.xlsx')
Image.new('RGB', (120, 80), '#35654e').save(out / 'sample.png')
(out / 'data.csv').write_text('Period,Energy\nMorning,10\nNoon,20\n')
# Minimal text PDF with correct xref offsets, no PDF-generation dependency needed.
stream = f'BT /F1 12 Tf 40 700 Td ({text}) Tj ET'.encode()
objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>', b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>', b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>', b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
result = b'%PDF-1.4\n'; offsets = [0]
for i, obj in enumerate(objects, 1):
    offsets.append(len(result)); result += f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
xref = len(result)
result += b'xref\n0 6\n0000000000 65535 f \n'+b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
result += f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
(out / 'notes.pdf').write_bytes(result)
